"""Offline audit regression tests. Failures document unfixed defects.
Run: .venv/bin/python -m unittest discover -s audit -p 'test_audit.py' -v
"""
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock
import numpy as np
from bookrag.config import Config
from bookrag.schemas import Chunk, ScoredChunk, Page
from bookrag.index.store import Store
from bookrag.ingest.chunker import chunk_book, count_tokens
from bookrag.ingest.loaders import load_docx, book_id_for, strip_page_furniture
from bookrag.ingest.structure import annotate_structure
from bookrag.retrieve.pipeline import Retriever, RetrievalResult
from bookrag.chat.engine import ChatEngine
from bookrag.llm.client import OllamaClient, parse_json
from bookrag.llm.router import FallbackClient


def chunk(book='a', ordinal=1, text='An ordinary passage about heat energy.'):
    return Chunk(f'{book}::{ordinal:05d}', text, text, book, book, page_start=ordinal, page_end=ordinal, ordinal=ordinal)


def line(text, chapter='Chapter 1', page=1, heading=False):
    return dict(text=text, chapter=chapter, section='', page=page, is_heading=heading)


class Audit(unittest.TestCase):
    def test_store_roundtrip(self):
        with tempfile.TemporaryDirectory() as d:
            s=Store(Path(d)); s.save([chunk(),chunk('b')],np.eye(2),{'books':[{'book_id':'a'},{'book_id':'b'}]})
            s=Store(Path(d)).load()
            self.assertEqual(s.dense_search(np.array([1,0]),1),[(0,1.0)])
            self.assertEqual(s.by_id('b::00001').book_id,'b')

    def test_unknown_book_filter_returns_nothing(self):
        with tempfile.TemporaryDirectory() as d:
            s=Store(Path(d)); s.save([chunk()],np.ones((1,2)),{'books':[{'book_id':'a'}]})
            self.assertEqual(s.dense_search(np.ones(2),10,['missing']),[])

    def test_repeated_store_save_refreshes_bm25(self):
        with tempfile.TemporaryDirectory() as d:
            s=Store(Path(d)); s.save([chunk(text='alpha'),chunk('b',text='beta')],np.eye(2),{'books':[]})
            s.bm25_search('alpha',2)
            s.save([chunk(text='gamma')],np.ones((1,2)),{'books':[]})
            self.assertEqual(len(s.bm25.get_scores(['gamma'])),1)

    def test_chunk_hard_token_limit(self):
        chunks=chunk_book([line('entropy '*2000)],'a','A',chunk_tokens=100,overlap_tokens=10,min_tokens=5)
        self.assertTrue(all(c.token_count<=110 for c in chunks),[c.token_count for c in chunks])

    def test_leading_fragment_preserves_chapter(self):
        chunks=chunk_book([line('INTRODUCTION',page=1),line('Heat is energy. '*100,'Chapter 2',page=2)],'a','A',min_tokens=60)
        intro=next(c for c in chunks if 'INTRODUCTION' in c.text)
        self.assertEqual((intro.chapter,intro.page_start),('Chapter 1',1))

    def test_page_citations_track_split_pages(self):
        lines=[line(('First page discusses energy. '*80),page=1),line(('Second page discusses entropy. '*80),page=2)]
        chunks=chunk_book(lines,'a','A',chunk_tokens=100,overlap_tokens=0,min_tokens=5)
        self.assertEqual((chunks[0].page_start,chunks[0].page_end),(1,1))

    def test_distinct_files_have_distinct_ids(self):
        with tempfile.TemporaryDirectory() as d:
            a=Path(d)/'a';b=Path(d)/'b';a.mkdir();b.mkdir()
            p=a/'book.txt';q=b/'book.txt';p.write_text('AAAA');q.write_text('BBBB')
            self.assertNotEqual(book_id_for(p),book_id_for(q))

    def test_docx_table_content_is_loaded(self):
        import docx
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'table.docx';doc=docx.Document();doc.add_paragraph('Ordinary body.');table=doc.add_table(rows=1,cols=1);table.cell(0,0).text='CRITICAL_TABLE_CONTENT';doc.save(p)
            pages,_=load_docx(p)
            self.assertIn('CRITICAL_TABLE_CONTENT',' '.join(p.text for p in pages))

    def test_header_removal_preserves_body_occurrence(self):
        pages=[]
        for i in range(5):
            blocks=[dict(text='Repeated Header',y=0,size=11),dict(text='Unique body '+str(i),y=100,size=11),dict(text='Bottom '+str(i),y=200,size=11)]
            if i==0:blocks.insert(2,dict(text='Repeated Header',y=150,size=11))
            pages.append(Page(i+1,'',blocks))
        cleaned=strip_page_furniture(pages)
        self.assertTrue(any(b['text']=='Repeated Header' and b['y']==150 for b in cleaned[0].blocks))

    def test_markdown_h2_does_not_become_chapter(self):
        from bookrag.ingest.loaders import load_text
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'book.md';p.write_text('# Chapter One\n'+('Body text here. '*30)+'\n## Heat Capacity\n'+('More body text. '*30))
            pages,meta=load_text(p);lines=annotate_structure(pages,meta)
            sub=next(l for l in lines if l['text']=='Heat Capacity')
            self.assertEqual(sub['chapter'],'Chapter One')

    def test_context_contains_top_hit_after_expansion(self):
        with tempfile.TemporaryDirectory() as d:
            s=Store(Path(d));cs=[chunk(ordinal=i,text=('background '*200 if i<3 else 'CRITICAL ANSWER')) for i in range(1,4)]
            s.save(cs,np.ones((3,2)),{'books':[{'book_id':'a'}]})
            r=object.__new__(Retriever);r.store=s;r.cfg=Config({'retrieval':{'neighbor_window':1,'max_context_tokens':100}})
            merged=r._expand_neighbors([ScoredChunk(cs[2],1,rerank_score=.99)])
            context,_=r.build_context(merged)
            self.assertIn('CRITICAL ANSWER',context)

    def test_empty_context_does_not_call_llm(self):
        rr=RetrievalResult('q',['q'],[],['a::00001'],'',[],.99,True)
        r=SimpleNamespace(reranker=True,retrieve=Mock(return_value=rr),arbiter=SimpleNamespace(before=Mock()))
        llm=Mock();llm.chat.return_value='Unsupported.'
        engine=ChatEngine(Config(),retriever=r,llm=llm);engine.ask('Explain entropy')
        llm.chat.assert_not_called()

    def test_invalid_citation_is_not_accepted(self):
        rr=RetrievalResult('q',['q'],[],[], 'source', ['Book A p.1'], .9,True)
        engine=ChatEngine(Config(),retriever=SimpleNamespace(),llm=Mock())
        answer=engine._finalize('q','q','Heat is energy [1]. Invented statement [99].',rr)
        self.assertFalse(answer.grounded)

    def test_uncited_answer_refused(self):
        rr=RetrievalResult('q',['q'],[],[],'source',['Book A p.1'],.9,True)
        engine=ChatEngine(Config(),retriever=SimpleNamespace(),llm=Mock())
        self.assertFalse(engine._finalize('q','q','Uncited text',rr).grounded)

    def test_refusal_stream_contract(self):
        rr=RetrievalResult('q',['q'],[],[],'',[],0,False)
        r=SimpleNamespace(reranker=True,retrieve=Mock(return_value=rr))
        engine=ChatEngine(Config(),retriever=r,llm=Mock())
        self.assertIn("can't answer",''.join(engine.ask('q',stream=True)))
        self.assertFalse(engine.last_answer.grounded)

    def test_ollama_health_checks_exact_tag(self):
        llm=OllamaClient(primary='qwen3:8b');llm.available_models=lambda:['qwen3:4b']
        self.assertFalse(llm.health_check()['primary_ok'])

    def test_router_respects_model_unavailable(self):
        p=SimpleNamespace(health_check=lambda:dict(primary_ok=False,fast_ok=False,models=['other']))
        self.assertFalse(FallbackClient(p,None).health_check()['primary_ok'])

    def test_fallback_before_first_token(self):
        p=Mock();p.stream_chat.side_effect=RuntimeError('offline');f=Mock();f.stream_chat.return_value=iter(['fallback'])
        self.assertEqual(list(FallbackClient(p,f).stream_chat([])),['fallback'])

    def test_no_fallback_after_first_token(self):
        def stream():
            yield 'start'
            raise RuntimeError('disconnect')
        p=Mock();p.stream_chat.return_value=stream();f=Mock();r=FallbackClient(p,f)
        with self.assertRaises(RuntimeError):list(r.stream_chat([]))
        f.stream_chat.assert_not_called()

    def test_json_repair(self):
        self.assertEqual(parse_json('```json\n{"x":1}\n```'),{'x':1})
        self.assertIsNone(parse_json('not JSON'))

if __name__=='__main__':unittest.main(verbosity=2)
