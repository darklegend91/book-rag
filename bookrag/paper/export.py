"""Render a generated Paper to Markdown, DOCX or JSON."""
from __future__ import annotations

import json
import re
from datetime import date
from pathlib import Path

from bookrag.paper.generator import Paper

# Fields that give the answer away. Dropped from every format when the export
# is meant for students.
_KEY_FIELDS = ("answer", "correct_option", "verify_notes")

_WINDOWS_RESERVED = {"con", "prn", "aux", "nul", *(f"com{i}" for i in range(1, 10)),
                     *(f"lpt{i}" for i in range(1, 10))}


def safe_stem(text: str, fallback: str = "paper", max_len: int = 80) -> str:
    """A filename stem that stays inside the export directory on every OS.

    A paper title is user text: "Physics I/II" used verbatim became a nested
    path that didn't exist, and "../x" would escape the export dir. Anything
    but letters, digits, '-' and '_' becomes '_'.
    """
    stem = re.sub(r"[^\w\-]+", "_", text.strip())
    stem = re.sub(r"_+", "_", stem).strip("._-")[:max_len].strip("._-")
    if stem.casefold() in _WINDOWS_RESERVED:
        stem = f"{stem}_"
    return stem or fallback


def _section_instructions(paper: Paper, name: str) -> str:
    """The blueprint's per-section instruction line, e.g. 'Answer ANY TWO.'"""
    spec = paper.section_spec(name)
    return (spec.instructions or "") if spec else ""


def _grouped(paper: Paper) -> list[tuple[str, list]]:
    order: list[str] = []
    buckets: dict[str, list] = {}
    for q in paper.questions:
        if q.section not in buckets:
            buckets[q.section] = []
            order.append(q.section)
        buckets[q.section].append(q)
    return [(name, buckets[name]) for name in order]


def _marks_note(paper: Paper, name: str, questions: list) -> str:
    """'(answer 1 of 3 questions × 2 marks = 2 marks)' -- shared by every format,
    so a student reading the DOCX sees the same attempt rule as the Markdown."""
    marks = paper.section_marks(name)
    spec = paper.section_spec(name)
    need = min(spec.answer_count, len(questions)) if spec else len(questions)
    if need < len(questions):
        return (f"(answer {need} of {len(questions)} questions × "
                f"{questions[0].marks} marks = {marks} marks)")
    return f"({len(questions)} questions × {questions[0].marks} marks = {marks} marks)"


def to_markdown(paper: Paper, include_key: bool = True,
                include_citations: bool = True) -> str:
    bp = paper.blueprint
    L: list[str] = [f"# {bp.title}"]
    if bp.subject:
        L.append(f"**Subject:** {bp.subject}  ")
    L.append(f"**Time:** {bp.duration}  •  **Maximum Marks:** {paper.total_marks}")
    L.append("")
    instructions = bp.general_instructions or [
        "Answer all questions unless otherwise stated.",
        "Marks are indicated against each question.",
        "Draw diagrams where appropriate.",
    ]
    L.append("**Instructions:**")
    L += [f"{i}. {line}" for i, line in enumerate(instructions, 1)]
    L.append("\n---\n")

    for name, questions in _grouped(paper):
        L.append(f"## {name}")
        instr = _section_instructions(paper, name)
        if instr:
            L.append(f"*{instr}*  ")
        L.append(f"*{_marks_note(paper, name, questions)}*\n")
        for q in questions:
            L.append(f"**Q{q.number}.** {q.text}  **[{q.marks}]**")
            if q.options:
                L += [f"   - ({chr(97+i)}) {o}" for i, o in enumerate(q.options)]
            L.append("")
        L.append("")

    if include_key:
        L.append("\n---\n\n# Answer Key\n")
        for name, questions in _grouped(paper):
            L.append(f"## {name}\n")
            for q in questions:
                L.append(f"**Q{q.number}.** {q.text}\n")
                if q.options and q.correct_option:
                    L.append(f"**Correct option:** {q.correct_option}\n")
                L.append(f"{q.answer or '_(no model answer generated)_'}\n")
                meta = [f"Bloom: {q.bloom}", f"Topic: {q.topic}",
                        f"verified {q.verify_confidence:.2f}" if q.verified else "unverified"]
                L.append(f"<sub>{' • '.join(meta)}</sub>\n")
                if include_citations and q.citations:
                    L.append("<sub>Source: " + "; ".join(q.citations[:3]) + "</sub>\n")
            L.append("")

    return "\n".join(L)


def to_docx(paper: Paper, path: Path, include_key: bool = True) -> Path:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    bp = paper.blueprint
    doc = docx.Document()
    title = doc.add_heading(bp.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    head.add_run(f"Time: {bp.duration}          Maximum Marks: {paper.total_marks}").bold = True

    doc.add_paragraph("Instructions:").runs[0].bold = True
    for line in (bp.general_instructions or [
        "Answer all questions unless otherwise stated.",
        "Marks are indicated against each question.",
    ]):
        doc.add_paragraph(line, style="List Number")

    for name, questions in _grouped(paper):
        doc.add_heading(name, level=1)
        instr = _section_instructions(paper, name)
        if instr:
            note = doc.add_paragraph(instr)
            note.runs[0].italic = True
        marks_note = doc.add_paragraph(_marks_note(paper, name, questions))
        marks_note.runs[0].italic = True
        marks_note.runs[0].font.size = Pt(9)
        for q in questions:
            p = doc.add_paragraph()
            p.add_run(f"Q{q.number}. ").bold = True
            p.add_run(q.text)
            marks = p.add_run(f"   [{q.marks}]")
            marks.bold = True
            marks.font.size = Pt(10)
            for i, opt in enumerate(q.options):
                doc.add_paragraph(f"({chr(97+i)}) {opt}", style="List Bullet")

    if include_key:
        doc.add_page_break()
        doc.add_heading("Answer Key", level=0)
        for name, questions in _grouped(paper):
            doc.add_heading(name, level=1)
            for q in questions:
                doc.add_paragraph(f"Q{q.number}. {q.text}").runs[0].bold = True
                if q.correct_option:
                    doc.add_paragraph(f"Correct option: {q.correct_option}")
                doc.add_paragraph(q.answer or "(no model answer generated)")
                if not q.verified:
                    doc.add_paragraph("(unverified)").runs[0].italic = True
                if q.citations:
                    src = doc.add_paragraph("Source: " + "; ".join(q.citations[:3]))
                    src.runs[0].font.size = Pt(8)
                    src.runs[0].italic = True

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def to_json(paper: Paper, include_key: bool = True) -> str:
    questions = [q.to_dict() for q in paper.questions]
    report = {
        "requested": paper.report.requested,
        "accepted": paper.report.accepted,
        "rejected": paper.report.rejected,
        "retried": paper.report.retried,
        "duplicates": paper.report.duplicates,
        "cancelled": paper.report.cancelled,
        "acceptance_rate": round(paper.report.acceptance_rate, 3),
        "ungrounded_topics": paper.report.ungrounded_topics,
        "rejections": paper.report.rejections,
    }
    if not include_key:
        for q in questions:
            for k in _KEY_FIELDS:
                q.pop(k, None)
        # Moderator feedback quotes drafts and explains why they were wrong,
        # which is as good as a key for questions that survived a retry.
        report.pop("rejections")
    return json.dumps({
        "blueprint": paper.blueprint.to_dict(),
        "total_marks": paper.total_marks,
        "generated_on": date.today().isoformat(),
        "includes_answer_key": include_key,
        "report": report,
        "questions": questions,
    }, indent=2, ensure_ascii=False)


def export_paper(paper: Paper, out_dir: Path, stem: str,
                 formats: tuple[str, ...] = ("md", "docx", "json"),
                 include_key: bool = True) -> list[Path]:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = safe_stem(stem)
    written: list[Path] = []
    if "md" in formats:
        p = out_dir / f"{stem}.md"
        p.write_text(to_markdown(paper, include_key), encoding="utf-8")
        written.append(p)
    if "json" in formats:
        p = out_dir / f"{stem}.json"
        p.write_text(to_json(paper, include_key), encoding="utf-8")
        written.append(p)
    if "docx" in formats:
        try:
            written.append(to_docx(paper, out_dir / f"{stem}.docx", include_key))
        except Exception as exc:
            print(f"  ! DOCX export skipped: {exc}")
    return written
